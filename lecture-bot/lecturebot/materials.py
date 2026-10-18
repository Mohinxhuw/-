"""Lecture files the student sends along with the recording: PDF and PowerPoint slides.

Claude reads PDFs natively (text and page images). PowerPoint files are turned into slide text plus
the pictures on the slides. English terms found in either are also given to the speech recognizer.
"""
import io
import re
from collections import Counter
from dataclasses import dataclass, field

PDF = "application/pdf"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
IMAGE_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp"}
MAX_PDF_PAGES = 100  # Claude's limit for one PDF


@dataclass
class Material:
    name: str
    pdf: bytes | None = None                      # sent to Claude as a document
    text: str = ""                                # slide text (PowerPoint) or extracted text (PDF)
    images: list[tuple[bytes, str]] = field(default_factory=list)


def kind_of(file_name: str, mime: str | None) -> str | None:
    name = (file_name or "").lower()
    if mime == PDF or name.endswith(".pdf"):
        return "pdf"
    if mime == PPTX or name.endswith(".pptx"):
        return "pptx"
    return None


def load(data: bytes, file_name: str, mime: str | None) -> Material:
    kind = kind_of(file_name, mime)
    if kind == "pdf":
        return _load_pdf(data, file_name)
    if kind == "pptx":
        return _load_pptx(data, file_name)
    raise ValueError(f"unsupported file: {file_name}")


def _load_pdf(data: bytes, name: str) -> Material:
    import pypdfium2

    doc = pypdfium2.PdfDocument(data)
    pages = len(doc)
    text = "\n".join(doc[i].get_textpage().get_text_range() for i in range(min(pages, MAX_PDF_PAGES)))
    if pages > MAX_PDF_PAGES:  # keep the first pages only
        out = pypdfium2.PdfDocument.new()
        out.import_pages(doc, list(range(MAX_PDF_PAGES)))
        buf = io.BytesIO()
        out.save(buf)
        data = buf.getvalue()
    return Material(name=name, pdf=data, text=text)


def _load_pptx(data: bytes, name: str, max_images: int = 12) -> Material:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(io.BytesIO(data))
    lines, images = [], []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if t:
                    texts.append(t)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text for c in row.cells))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and len(images) < max_images:
                img = shape.image
                mime = img.content_type
                if mime in IMAGE_TYPES and len(img.blob) > 20_000:  # skip logos and icons
                    images.append((img.blob, mime))
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if texts or notes.strip():
            lines.append(f"[شريحة {i}]\n" + "\n".join(texts) + (f"\nملاحظات: {notes}" if notes.strip() else ""))
    return Material(name=name, text="\n\n".join(lines), images=images)


_STOP = {"the", "and", "for", "with", "from", "this", "that", "are", "was", "were", "has", "have", "not",
         "but", "can", "will", "into", "its", "our", "your", "you", "all", "any", "how", "what", "which",
         "when", "where", "who", "why", "also", "more", "most", "such", "than", "then", "there", "these",
         "they", "them", "each", "other", "some", "only", "very", "use", "used", "using", "page", "slide",
         "chapter", "lecture", "figure", "table", "example", "part", "one", "two", "three"}


def key_terms(texts: list[str], limit: int = 25) -> list[str]:
    """English terms that appear in the slides, most frequent first, for the speech recognizer."""
    counts: Counter[str] = Counter()
    display: dict[str, str] = {}
    for text in texts:
        for w in re.findall(r"[A-Za-z][A-Za-z\-]{2,}", text or ""):
            k = w.lower()
            if k in _STOP:
                continue
            counts[k] += 1
            # Terms are written capitalised on slides (Mitochondria, ATP); ordinary words mostly are not.
            if w[0].isupper():
                display.setdefault(k, w if not w.isupper() or len(w) <= 5 else w.capitalize())
    ranked = [k for k, _ in counts.most_common() if k in display]
    return [display[k] for k in ranked[:limit]]
