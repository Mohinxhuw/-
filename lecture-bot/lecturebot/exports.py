"""Editable copies of a lecture's notes: Word (right-to-left) and Markdown (imports cleanly into Notion),"""
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import mindmap
from .render import LABELS, _letter

NAVY = RGBColor(0x14, 0x3F, 0x62)
LETTERS = LABELS["ar"]["letters"]


# ---------- Word ----------

def _bidi(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:bidi")
    el.set(qn("w:val"), "1")
    ppr.append(el)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _run(paragraph, text: str, bold: bool = False, size: float | None = None, color=None):
    r = paragraph.add_run(text)
    if bold:
        r.bold = True
        r._r.get_or_add_rPr().append(OxmlElement("w:bCs"))  # bold for Arabic (complex script) text
    if size:
        r.font.size = Pt(size)
        sz = OxmlElement("w:szCs")
        sz.set(qn("w:val"), str(int(size * 2)))
        r._r.get_or_add_rPr().append(sz)
    if color is not None:
        r.font.color.rgb = color
    return r


class _Doc:
    def __init__(self):
        self.doc = Document()
        style = self.doc.styles["Normal"]
        style.font.name = "Arial"
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), "Arial")
        style.font.size = Pt(12)
        for section in self.doc.sections:
            section.left_margin = section.right_margin = Cm(2)

    def p(self, text: str = "", bold: bool = False, size: float | None = None, color=None, style: str | None = None):
        par = self.doc.add_paragraph(style=style)
        _bidi(par)
        if text:
            _run(par, text, bold, size, color)
        return par

    def h(self, text: str, level: int = 1):
        return self.p(text, bold=True, size={0: 20, 1: 16, 2: 13}[level], color=NAVY)

    def bullet(self, text: str):
        return self.p(text, style="List Bullet")

    def table(self, header: list[str], rows: list[list[str]]):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t._tbl.tblPr.append(OxmlElement("w:bidiVisual"))  # columns from right to left
        for cell, text in zip(t.rows[0].cells, header):
            cell.text = ""
            _bidi(cell.paragraphs[0])
            _run(cell.paragraphs[0], text, bold=True)
        for row in rows:
            cells = t.add_row().cells
            for cell, text in zip(cells, row):
                cell.text = ""
                _bidi(cell.paragraphs[0])
                _run(cell.paragraphs[0], str(text))
        self.p()
        return t

    def bytes(self) -> bytes:
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


def _with_time(text: str, time: str) -> str:
    return f"{text} ({time})" if time else text


def lecture_docx(lec: dict) -> bytes:
    d = _Doc()
    d.h(lec.get("title") or "ملخص المحاضرة", 0)
    if lec.get("subject"):
        d.p(lec["subject"], color=RGBColor(0x5B, 0x68, 0x76))
    if lec.get("overview"):
        d.p(lec["overview"])
    d.h("الملخص")
    for sec in lec.get("sections", []):
        d.h(_with_time(sec.get("heading", ""), sec.get("time", "")), 2)
        for pt in sec.get("points", []):
            d.bullet(_with_time(pt.get("text", ""), pt.get("time", "")))
    if lec.get("divisions"):
        d.h("التقسيمات")
        for div in lec["divisions"]:
            d.h(_with_time(div.get("title", ""), div.get("time", "")), 2)
            png = mindmap.render_png(div, width_px=1800)
            if png:
                d.doc.add_picture(io.BytesIO(png), width=Cm(17))
    for c in lec.get("comparisons", []):
        d.h(c["title"], 2)
        d.table(c["headers"], c["rows"])
    if lec.get("glossary"):
        d.h("المصطلحات")
        d.table(["المصطلح", "بالعربية", "المعنى"],
                [[g.get("term", ""), g.get("arabic", ""), g.get("explanation", "")] for g in lec["glossary"]])
    if lec.get("mcq"):
        d.h("أسئلة الاختيار من متعدد")
        for i, q in enumerate(lec["mcq"], 1):
            d.p(f"{i}. {q['question']}", bold=True)
            for j, o in enumerate(q["options"]):
                d.p(f"({_letter(j)}) {o}")
        d.h("الإجابات", 2)
        for i, q in enumerate(lec["mcq"], 1):
            d.p(f"{i}. ({_letter(q['answer_index'])}) {q['options'][q['answer_index']]} — {q.get('explanation', '')}")
    if lec.get("true_false"):
        d.h("أسئلة صح أو خطأ")
        for i, t in enumerate(lec["true_false"], 1):
            d.p(f"{i}. {t['statement']}  (      )")
        d.h("الإجابات", 2)
        for i, t in enumerate(lec["true_false"], 1):
            d.p(f"{i}. {'صح' if t['answer'] else 'خطأ'} — {t.get('explanation', '')}")
    if lec.get("essay"):
        d.h("أسئلة مقالية")
        for i, q in enumerate(lec["essay"], 1):
            d.p(f"{i}. {q['question']}", bold=True)
            d.p(f"نموذج الإجابة: {q.get('model_answer', '')}")
    if lec.get("flashcards"):
        d.h("بطاقات المراجعة")
        d.table(["السؤال", "الجواب"], [[c["front"], c["back"]] for c in lec["flashcards"]])
    if lec.get("transcript"):
        d.h("النص الكامل")
        for para in lec["transcript"].split("\n\n"):
            d.p(para, size=10)
    return d.bytes()


# ---------- Markdown (Notion) ----------

def _cell(text) -> str:
    return str(text).replace("|", "/").replace("\n", " ")


def lecture_markdown(lec: dict) -> str:
    out = [f"# {lec.get('title') or 'ملخص المحاضرة'}", ""]
    if lec.get("subject"):
        out += [f"*{lec['subject']}*", ""]
    if lec.get("overview"):
        out += [f"> {lec['overview']}", ""]
    out += ["## الملخص", ""]
    for sec in lec.get("sections", []):
        out += [f"### {_with_time(sec.get('heading', ''), sec.get('time', ''))}", ""]
        out += [f"- {_with_time(p.get('text', ''), p.get('time', ''))}" for p in sec.get("points", [])] + [""]
    if lec.get("divisions"):
        out += ["## التقسيمات", ""]
        for div in lec["divisions"]:
            out += [f"### {_with_time(div.get('title', ''), div.get('time', ''))}", "", f"- **{div.get('center', '')}**"]
            for b in div.get("branches", []):
                out.append(f"  - **{b['label']}**")
                out += [f"    - {c}" for c in b.get("children", [])]
            out.append("")
    for c in lec.get("comparisons", []):
        out += [f"### {c['title']}", "", "| " + " | ".join(_cell(h) for h in c["headers"]) + " |",
                "|" + "---|" * len(c["headers"])]
        out += ["| " + " | ".join(_cell(x) for x in r) + " |" for r in c["rows"]] + [""]
    if lec.get("glossary"):
        out += ["## المصطلحات", "", "| المصطلح | بالعربية | المعنى |", "|---|---|---|"]
        out += [f"| {_cell(g.get('term', ''))} | {_cell(g.get('arabic', ''))} | {_cell(g.get('explanation', ''))} |"
                for g in lec["glossary"]] + [""]
    if lec.get("mcq"):
        out += ["## أسئلة الاختيار من متعدد", ""]
        for i, q in enumerate(lec["mcq"], 1):
            out += [f"**{i}. {q['question']}**", ""] + [f"- ({_letter(j)}) {o}" for j, o in enumerate(q["options"])]
            out += ["", f"الإجابة: ({_letter(q['answer_index'])}) — {q.get('explanation', '')}", ""]
    if lec.get("true_false"):
        out += ["## أسئلة صح أو خطأ", ""]
        out += [f"{i}. {t['statement']} — **{'صح' if t['answer'] else 'خطأ'}**: {t.get('explanation', '')}"
                for i, t in enumerate(lec["true_false"], 1)] + [""]
    if lec.get("essay"):
        out += ["## أسئلة مقالية", ""]
        for i, q in enumerate(lec["essay"], 1):
            out += [f"**{i}. {q['question']}**", "", f"نموذج الإجابة: {q.get('model_answer', '')}", ""]
    if lec.get("flashcards"):
        out += ["## بطاقات المراجعة", "", "| السؤال | الجواب |", "|---|---|"]
        out += [f"| {_cell(c['front'])} | {_cell(c['back'])} |" for c in lec["flashcards"]] + [""]
    if lec.get("transcript"):
        out += ["## النص الكامل", "", lec["transcript"], ""]
    return "\n".join(out)
